# -*- coding: utf-8 -*-
"""
SIEAGE - Importador del libro PRINCIPAL-2026.xlsx a la base `sieage`.

Lee las 11 hojas de grupo, deduplica estudiantes/acudientes/barrios,
desdobla las 12 columnas dobles de anio en filas de `matriculas`, y
genera:

  docs/bd/datos-importados.sql   -> los INSERT, listos para cargar
  docs/bd/reporte-importacion.txt-> resumen legible
  docs/bd/pendiente-*.csv        -> filas que hay que revisar a mano

No toca la base: solo produce el .sql. La carga es un paso aparte.

Uso:  python docs/bd/importar_excel.py "<ruta del .xlsx>"
"""
import openpyxl, collections, csv, io, os, re, sys, unicodedata

sys.stdout.reconfigure(encoding='utf-8')

XLSX = sys.argv[1] if len(sys.argv) > 1 else \
    r"C:\Users\Kechavarro\Downloads\PRINCIPAL-2026 (4) (1).xlsx"
OUT = os.path.dirname(os.path.abspath(__file__))
HOJAS = ['10-1', '10-2', '10-3', '10-4', '10-5', '10-6',
         '11-1', '11-2', '11-3', '11-4', '11-5']

# --- IDs de catalogo, tal como quedaron sembrados en la base ---------------
SEDE = {'P': 1, 'LF': 2, 'CP': 3, 'RP': 4, 'PT': 5}
GRADO = {n: n + 1 for n in range(0, 12)}          # grado 0 -> id 1 ... 11 -> 12
MODALIDAD = {'Asist. Admin.': 1, 'Ebanistería': 2, 'Electricidad': 3,
             'Electrónica': 4}
PARENTESCO = {'Madre': 1, 'Padre': 2, 'Hermana(o)': 3, 'Abuela(o)': 4,
              'Tia(o)': 5, 'Madrastra': 6, 'Padrastro': 7, 'Prima(o)': 8,
              'Otro': 9}
ANIO = {a: i + 1 for i, a in enumerate(range(2015, 2027))}
ANIO_ACTUAL = 2026

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5,
         'junio': 6, 'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9,
         'octubre': 10, 'noviembre': 11, 'diciembre': 12}

# Columna del primer par de anio (I=9) y su codigo de sede (J=10)
COL_ANIO_INICIO = 9
COL_BOLETIN = range(41, 51)          # AO..AX = boletin 1º..10º

TIPO_DOC_OK = {'R.C.', 'T.I.', 'C.C.', 'C.E.', 'P.P.T.', 'N.U.I.P.', 'N.E.S.'}
TIPO_DOC_ACU = {'C.C.', 'C.E.', 'P.P.T.', 'PAS', 'N.I.T.'}

PARENTESCO_ALIAS = {
    'prima': 'Prima(o)', 'primo': 'Prima(o)', 'prima(o)': 'Prima(o)',
    'hermana': 'Hermana(o)', 'hermano': 'Hermana(o)',
    'hermana(o)': 'Hermana(o)', 'abuela': 'Abuela(o)', 'abuelo': 'Abuela(o)',
    'abuela(o)': 'Abuela(o)', 'tia': 'Tia(o)', 'tio': 'Tia(o)',
    'tia(o)': 'Tia(o)', 'madre': 'Madre', 'padre': 'Padre',
    'madrastra': 'Madrastra', 'padrastro': 'Padrastro',
}

incidencias = collections.defaultdict(list)


# ------------------------------------------------------------------ utils --
def txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r'\s+', ' ', str(v)).strip()


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s)
                   if unicodedata.category(c) != 'Mn').lower()


def sql(v):
    """Literal SQL seguro."""
    if v is None or v == '':
        return 'NULL'
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace('\\', '\\\\').replace("'", "''") + "'"


def fecha(d, m, a):
    """Arma YYYY-MM-DD desde tres celdas; el mes puede venir en texto."""
    d, a = txt(d), txt(a)
    m = txt(m)
    if not (d and m and a):
        return None
    if not m.isdigit():
        m = MESES.get(sin_tildes(m))
        if not m:
            return None
    try:
        d, m, a = int(d), int(m), int(a)
        if not (1 <= d <= 31 and 1 <= m <= 12 and 1900 < a < 2100):
            return None
        return '%04d-%02d-%02d' % (a, m, d)
    except ValueError:
        return None


def parse_grupo(valor):
    """'10-3' -> (10, 3) · '9º' -> (9, None) · 'No'/'' -> None (no cursó aquí)."""
    s = txt(valor)
    if not s or s.lower() == 'no':
        return None
    m = re.match(r'^(\d{1,2})\s*-\s*(\d{1,2})$', s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r'^(\d{1,2})\s*[ºo°]?$', s)
    if m:
        return int(m.group(1)), None
    return 'INVALIDO'


# ------------------------------------------------------------------- carga --
print('Leyendo %s ...' % os.path.basename(XLSX))
wb = openpyxl.load_workbook(XLSX, data_only=True)

estudiantes = {}          # (tipo, doc) -> dict
acudientes = {}           # (tipo, doc) -> dict
barrios = {}              # nombre normalizado -> id
grupos = {}               # (anio, sede_id, grado, numero, jornada) -> dict
matriculas = []
boletines_pendientes = []
sede_discrepancia = []
grupo_discrepancia = []

for hoja in HOJAS:
    ws = wb[hoja]
    grado_hoja, numero_hoja = (int(x) for x in hoja.split('-'))
    for r in range(3, ws.max_row + 1):
        C = lambda c: ws.cell(row=r, column=c).value

        nombre = txt(C(6))
        if not nombre:
            continue
        fila = '%s!fila %d' % (hoja, r)

        # ---- documento del estudiante ---------------------------------
        tipo = txt(C(39)) or 'T.I.'
        if tipo not in TIPO_DOC_OK:
            incidencias['tipo_doc_desconocido'].append(
                '%s · %s · tipo "%s" -> se usó T.I.' % (fila, nombre, tipo))
            tipo = 'T.I.'
        doc = txt(C(40))
        if not doc:
            incidencias['sin_documento'].append('%s · %s' % (fila, nombre))
            continue

        clave = (tipo, doc)
        if clave in estudiantes:
            incidencias['documento_duplicado'].append(
                '%s · %s · doc %s ya usado por "%s" -> fila omitida'
                % (fila, nombre, doc, estudiantes[clave]['nombre']))
            continue

        est_id = len(estudiantes) + 1
        estudiantes[clave] = {
            'id': est_id, 'tipo': tipo, 'doc': doc, 'nombre': nombre,
            'nacimiento': fecha(C(35), C(36), C(37)),
            'genero': txt(C(34)) if txt(C(34)) in ('F', 'M') else None,
            'foto': txt(C(51)).upper() == 'S',
        }
        if not estudiantes[clave]['nacimiento']:
            incidencias['sin_fecha_nacimiento'].append(
                '%s · %s' % (fila, nombre))

        # ---- acudiente -------------------------------------------------
        acu_doc = txt(C(53))
        acu_nombre = txt(C(52))
        acu_id = None
        if acu_doc and acu_nombre:
            kacu = ('C.C.', acu_doc)
            if kacu not in acudientes:
                barrio = txt(C(56))
                bid = None
                if barrio:
                    # La base usa utf8mb4_unicode_ci: "Alfonso López" y
                    # "Alfonso Lopez" son el MISMO barrio para el índice
                    # UNIQUE. Se agrupa sin tildes y se conserva la grafía
                    # más frecuente como nombre a mostrar.
                    bkey = sin_tildes(barrio)
                    if bkey not in barrios:
                        barrios[bkey] = {'id': len(barrios) + 1,
                                         'formas': collections.Counter()}
                    barrios[bkey]['formas'][barrio.title()] += 1
                    bid = barrios[bkey]['id']
                acudientes[kacu] = {
                    'id': len(acudientes) + 1, 'doc': acu_doc,
                    'nombre': acu_nombre, 'fijo': txt(C(57)) or None,
                    'cel': txt(C(58)) or None, 'dir': txt(C(55)) or None,
                    'barrio_id': bid, 'hijos': [],
                }
            acu = acudientes[kacu]
            par = PARENTESCO_ALIAS.get(sin_tildes(txt(C(54))), None)
            if par is None and txt(C(54)):
                incidencias['parentesco_normalizado'].append(
                    '%s · "%s" -> Otro' % (fila, txt(C(54))))
            acu['hijos'].append((est_id, PARENTESCO[par or 'Otro']))
        elif acu_nombre and not acu_doc:
            incidencias['acudiente_sin_cedula'].append(
                '%s · %s · acudiente "%s"' % (fila, nombre, acu_nombre))

        # ---- recorrido de los 12 anios ---------------------------------
        historial = {}
        for i, anio in enumerate(range(2015, 2027)):
            col_g = COL_ANIO_INICIO + i * 2
            g = parse_grupo(C(col_g))
            if g is None:
                continue
            if g == 'INVALIDO':
                incidencias['grupo_ilegible'].append(
                    '%s · %s · %d = "%s"' % (fila, nombre, anio, txt(C(col_g))))
                continue
            cod_sede = txt(C(col_g + 1)).upper()
            if anio == ANIO_ACTUAL:
                cod_sede = cod_sede or 'P'
            if cod_sede not in SEDE:
                incidencias['sede_desconocida'].append(
                    '%s · %s · %d código "%s" -> omitido'
                    % (fila, nombre, anio, cod_sede))
                continue
            historial[anio] = (g[0], g[1], SEDE[cod_sede])

        # El nombre de la hoja ES la lista oficial del grupo de 2026. La
        # columna AE quedó desactualizada en 27 filas (repite el grupo de
        # 2025, o dice solo "9°"/"10°"), así que manda la hoja y la
        # diferencia se reporta.
        ae_original = txt(C(31))
        sede_2026 = historial.get(ANIO_ACTUAL, (0, 0, SEDE['P']))[2]
        if ae_original != hoja:
            grupo_discrepancia.append(
                [hoja, nombre, doc, ae_original or '(vacío)', txt(C(29)), hoja])
        historial[ANIO_ACTUAL] = (grado_hoja, numero_hoja, sede_2026)

        # la columna G (SEDE) choca con el código del año 2026: se reporta
        sede_col = txt(C(7))
        if ANIO_ACTUAL in historial and sede_col == 'Los Farallones':
            sede_discrepancia.append(
                [hoja, nombre, doc, sede_col, 'Principal (código P del año 2026)'])

        # ---- una matrícula por año cursado ------------------------------
        anios = sorted(historial)
        for anio in anios:
            grado, numero, sede_id = historial[anio]
            actual = anio == ANIO_ACTUAL
            jornada = (txt(C(33)) or 'Mañana') if actual else 'Mañana'
            if jornada not in ('Mañana', 'Tarde', 'Única', 'Noche'):
                jornada = 'Mañana'

            gid = None
            if numero is not None:
                gkey = (anio, sede_id, grado, numero, jornada)
                if gkey not in grupos:
                    grupos[gkey] = {
                        'id': len(grupos) + 1, 'anio': anio, 'sede': sede_id,
                        'grado': grado, 'numero': numero,
                        'codigo': '%d-%d' % (grado, numero),
                        'jornada': jornada,
                        'cupos': 34 if actual else 34,
                    }
                gid = grupos[gkey]['id']
            else:
                incidencias['grupo_sin_numero'].append(
                    '%s · %s · %d grado %d sin número de grupo'
                    % (fila, nombre, anio, grado))

            # condición: nuevo / antiguo / repitente
            prev = historial.get(anio - 1)
            if prev is None:
                cond = 'nuevo'
            elif prev[0] == grado:
                cond = 'repitente'
            else:
                cond = 'antiguo'

            # resultado: se deduce del año siguiente
            sig = historial.get(anio + 1)
            resultado = None
            if sig:
                resultado = 'promovido' if sig[0] > grado else 'reprobado'

            # estado (solo el año en curso tiene marca en la columna A)
            estado = 'activo'
            if actual:
                marca = txt(C(1)).upper()
                if marca == 'R':
                    estado = 'retirado'
                elif marca != '1':
                    estado = 'cancelado'
                    incidencias['sin_marca_matricula'].append(
                        '%s · %s · columna A vacía -> estado "cancelado"'
                        % (fila, nombre))

            modalidad = None
            if actual:
                mod = txt(C(8))
                modalidad = MODALIDAD.get(mod)
                if mod and modalidad is None:
                    incidencias['modalidad_desconocida'].append(
                        '%s · "%s"' % (fila, mod))

            matriculas.append({
                'id': len(matriculas) + 1, 'est': est_id,
                'anio': ANIO[anio], 'grado': GRADO[grado], 'grupo': gid,
                'sede': sede_id, 'modalidad': modalidad,
                'jornada': jornada if actual else None,
                'fecha': fecha(C(3), C(4), C(5)) if actual else None,
                'cond': cond, 'estado': estado, 'resultado': resultado,
                'orden': int(txt(C(2))) if actual and txt(C(2)).isdigit() else None,
                'historico': not actual,
                'obs': (txt(C(59)) or None) if actual else None,
            })

        # ---- boletines: semántica sin confirmar, se guardan aparte ------
        bol = {}
        for i, c in enumerate(COL_BOLETIN, start=1):
            v = txt(C(c))
            if v:
                bol[i] = v
        if bol:
            boletines_pendientes.append(
                [hoja, nombre, doc] + [bol.get(i, '') for i in range(1, 11)])

print('  estudiantes  : %d' % len(estudiantes))
print('  acudientes   : %d' % len(acudientes))
print('  barrios      : %d' % len(barrios))
print('  grupos       : %d' % len(grupos))
print('  matrículas   : %d' % len(matriculas))

# ------------------------------------------------------------------- SQL ---
out = io.open(os.path.join(OUT, 'datos-importados.sql'), 'w',
              encoding='utf-8', newline='\n')
w = out.write
w('-- Generado por docs/bd/importar_excel.py desde %s\n'
  % os.path.basename(XLSX))
w('-- Cárgalo sobre la base `sieage` YA creada con su esquema y catálogos.\n')
w('SET NAMES utf8mb4;\nSET FOREIGN_KEY_CHECKS = 1;\nSTART TRANSACTION;\n\n')

w('-- barrios ------------------------------------------------------------\n')
for b in sorted(barrios.values(), key=lambda x: x['id']):
    nombre = b['formas'].most_common(1)[0][0]
    if len(b['formas']) > 1:
        incidencias['barrio_unificado'].append(
            '%s <- %s' % (nombre, ', '.join(sorted(b['formas']))))
    w('INSERT INTO barrios (id, nombre) VALUES (%d, %s);\n' % (b['id'], sql(nombre)))

w('\n-- grupos ------------------------------------------------------------\n')
for g in sorted(grupos.values(), key=lambda x: x['id']):
    w('INSERT INTO grupos (id, anio_lectivo_id, sede_id, grado_id, numero, '
      'codigo, jornada, cupos_proyectados) VALUES (%d, %d, %d, %d, %d, %s, %s, %d);\n'
      % (g['id'], ANIO[g['anio']], g['sede'], GRADO[g['grado']], g['numero'],
         sql(g['codigo']), sql(g['jornada']), g['cupos']))

w('\n-- estudiantes -------------------------------------------------------\n')
for e in sorted(estudiantes.values(), key=lambda x: x['id']):
    w('INSERT INTO estudiantes (id, tipo_documento, numero_documento, '
      'nombre_completo, fecha_nacimiento, genero, tiene_foto) VALUES '
      '(%d, %s, %s, %s, %s, %s, %s);\n'
      % (e['id'], sql(e['tipo']), sql(e['doc']), sql(e['nombre']),
         sql(e['nacimiento']), sql(e['genero']), sql(e['foto'])))

w('\n-- acudientes --------------------------------------------------------\n')
for a in sorted(acudientes.values(), key=lambda x: x['id']):
    w('INSERT INTO acudientes (id, tipo_documento, numero_documento, '
      'nombre_completo, telefono_fijo, telefono_celular, direccion, barrio_id) '
      'VALUES (%d, %s, %s, %s, %s, %s, %s, %s);\n'
      % (a['id'], sql('C.C.'), sql(a['doc']), sql(a['nombre']),
         sql(a['fijo']), sql(a['cel']), sql(a['dir']), sql(a['barrio_id'])))

w('\n-- vínculo estudiante-acudiente --------------------------------------\n')
for a in sorted(acudientes.values(), key=lambda x: x['id']):
    for est_id, par_id in a['hijos']:
        w('INSERT INTO estudiante_acudiente (estudiante_id, acudiente_id, '
          'parentesco_id) VALUES (%d, %d, %d);\n' % (est_id, a['id'], par_id))

w('\n-- matrículas (una fila por año cursado) -----------------------------\n')
for m in matriculas:
    w('INSERT INTO matriculas (id, estudiante_id, anio_lectivo_id, grado_id, '
      'grupo_id, sede_id, modalidad_id, jornada, fecha_matricula, condicion, '
      'estado, resultado, numero_orden, es_historico, observaciones) VALUES '
      '(%d, %d, %d, %d, %s, %d, %s, %s, %s, %s, %s, %s, %s, %s, %s);\n'
      % (m['id'], m['est'], m['anio'], m['grado'], sql(m['grupo']), m['sede'],
         sql(m['modalidad']), sql(m['jornada']), sql(m['fecha']),
         sql(m['cond']), sql(m['estado']), sql(m['resultado']),
         sql(m['orden']), sql(m['historico']), sql(m['obs'])))

w('\nCOMMIT;\n')
out.close()

# --------------------------------------------------------------- reportes --
def csv_out(nombre, cabecera, filas):
    if not filas:
        return None
    ruta = os.path.join(OUT, nombre)
    with io.open(ruta, 'w', encoding='utf-8-sig', newline='') as f:
        cw = csv.writer(f, delimiter=';')
        cw.writerow(cabecera)
        cw.writerows(filas)
    return ruta

csv_out('pendiente-boletines.csv',
        ['hoja', 'estudiante', 'documento'] + ['boletin_%d' % i for i in range(1, 11)],
        boletines_pendientes)
csv_out('pendiente-grupo-desactualizado.csv',
        ['hoja', 'estudiante', 'documento', 'columna_GRUPO_2026',
         'columna_GRUPO_2025', 'grupo_importado'],
        grupo_discrepancia)
csv_out('pendiente-sede-discrepancia.csv',
        ['hoja', 'estudiante', 'documento', 'columna_SEDE', 'sede_importada'],
        sede_discrepancia)

rep = io.open(os.path.join(OUT, 'reporte-importacion.txt'), 'w',
              encoding='utf-8', newline='\n')
rep.write('REPORTE DE IMPORTACIÓN - %s\n' % os.path.basename(XLSX))
rep.write('=' * 70 + '\n\n')
rep.write('REGISTROS GENERADOS\n')
for k, v in [('estudiantes', len(estudiantes)), ('acudientes', len(acudientes)),
             ('barrios', len(barrios)), ('grupos', len(grupos)),
             ('matrículas', len(matriculas))]:
    rep.write('  %-14s %d\n' % (k, v))

rep.write('\nMATRÍCULAS POR AÑO\n')
por_anio = collections.Counter(m['anio'] for m in matriculas)
inv = {v: k for k, v in ANIO.items()}
for aid in sorted(por_anio):
    rep.write('  %d  %d\n' % (inv[aid], por_anio[aid]))

rep.write('\nESTADO DEL AÑO %d\n' % ANIO_ACTUAL)
est_cnt = collections.Counter(m['estado'] for m in matriculas
                              if m['anio'] == ANIO[ANIO_ACTUAL])
for k, v in est_cnt.most_common():
    rep.write('  %-12s %d\n' % (k, v))

rep.write('\nINCIDENCIAS\n')
if not incidencias:
    rep.write('  ninguna\n')
for k in sorted(incidencias):
    rep.write('\n  [%s]  %d caso(s)\n' % (k, len(incidencias[k])))
    for linea in incidencias[k][:25]:
        rep.write('    - %s\n' % linea)
    if len(incidencias[k]) > 25:
        rep.write('    ... y %d más\n' % (len(incidencias[k]) - 25))
rep.close()

print('\nGenerado:')
print('  docs/bd/datos-importados.sql')
print('  docs/bd/reporte-importacion.txt')
if boletines_pendientes:
    print('  docs/bd/pendiente-boletines.csv (%d filas)' % len(boletines_pendientes))
if sede_discrepancia:
    print('  docs/bd/pendiente-sede-discrepancia.csv (%d filas)' % len(sede_discrepancia))
