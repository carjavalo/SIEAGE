# -*- coding: utf-8 -*-
"""
SIEAGE - Importador multi-sede de los libros de matricula 2026.

Procesa los 5 libros (uno por sede) en una sola pasada, de modo que los
estudiantes, acudientes, barrios y grupos compartidos entre sedes queden
unificados con un solo id.

Reglas de negocio aplicadas (documentadas en 00-analisis-excel.md):

  * El nombre de la hoja ES el grupo oficial del ano 2026; la columna AE
    esta desactualizada en parte de las filas.
  * La sede del ano 2026 es la del libro, no la columna G (metadato viejo)
    ni la columna AF (que a veces apunta a otra sede).
  * Los anos 2015-2025 se leen de los 11 pares de columnas (grupo + codigo
    de sede) y producen una fila de `matriculas` por ano cursado.
  * La jornada del grupo se decide por mayoria de sus estudiantes; la del
    estudiante se guarda tal como viene. 'Por definir' -> NULL.
  * `modalidad` solo se lee en 10o y 11o (es la media tecnica).
  * `condicion` y `resultado` se deducen comparando el grado de cada ano
    con el del ano anterior y el siguiente.

Genera (no toca la base):
  docs/bd/datos-sedes.sql          -> los INSERT
  docs/bd/reporte-sedes.txt        -> resumen legible
  docs/bd/pendiente-*.csv          -> filas a revisar a mano

Uso:  python docs/bd/importar_sedes.py [carpeta_de_los_xlsx]
"""
import openpyxl, collections, csv, io, os, re, sys, unicodedata

sys.stdout.reconfigure(encoding='utf-8')

DESCARGAS = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\Kechavarro\Downloads"
OUT = os.path.dirname(os.path.abspath(__file__))

# Orden de proceso: la Principal primero para que conserve los ids bajos.
LIBROS = [
    ('P',  'PRINCIPAL-2026 (4) (1).xlsx'),
    ('LF', 'LOS FARALLONES-2026 (1).xlsx'),
    ('RP', 'RAFAEL POMBO-2026.xlsx'),
    ('PT', 'PURIFICACION TRUJILLO-2026 (1).xlsx'),
    ('CP', 'CENTRAL PROVIVIENDA-2026 (1).xlsx'),
]

# --- ids de catalogo, tal como estan sembrados en la base -------------------
SEDE = {'P': 1, 'LF': 2, 'CP': 3, 'RP': 4, 'PT': 5}
SEDE_NOMBRE = {'P': 'Principal', 'LF': 'Los Farallones',
               'CP': 'Central Provivienda', 'RP': 'Rafael Pombo',
               'PT': 'Purificación Trujillo'}
GRADO = {n: n + 1 for n in range(0, 12)}
MODALIDAD = {'Asist. Admin.': 1, 'Ebanistería': 2, 'Electricidad': 3,
             'Electrónica': 4}
PARENTESCO = {'Madre': 1, 'Padre': 2, 'Hermana(o)': 3, 'Abuela(o)': 4,
              'Tia(o)': 5, 'Madrastra': 6, 'Padrastro': 7, 'Prima(o)': 8,
              'Otro': 9}
ANIO = {a: i + 1 for i, a in enumerate(range(2015, 2027))}
ANIO_ACTUAL = 2026
GRADOS_MEDIA = (10, 11)

MESES = {'enero': 1, 'febrero': 2, 'marzo': 3, 'abril': 4, 'mayo': 5,
         'junio': 6, 'julio': 7, 'agosto': 8, 'septiembre': 9, 'setiembre': 9,
         'octubre': 10, 'noviembre': 11, 'diciembre': 12}

COL_ANIO_INICIO = 9                  # I = GRUPO 2015, J = su codigo de sede
COL_BOLETIN = range(41, 51)          # AO..AX = boletin 1o..10o
JORNADAS_VALIDAS = ('Mañana', 'Tarde', 'Única', 'Noche')
JORNADA_ALIAS = {
    'manana': 'Mañana', 'mañana': 'Mañana', 'tarde': 'Tarde',
    'trade': 'Tarde',                      # typo real en LOS FARALLONES
    'unica': 'Única', 'noche': 'Noche',
}
TIPO_DOC_OK = {'R.C.', 'T.I.', 'C.C.', 'C.E.', 'P.P.T.', 'N.U.I.P.', 'N.E.S.'}
PARENTESCO_ALIAS = {
    'prima': 'Prima(o)', 'primo': 'Prima(o)', 'prima(o)': 'Prima(o)',
    'hermana': 'Hermana(o)', 'hermano': 'Hermana(o)',
    'hermana(o)': 'Hermana(o)', 'abuela': 'Abuela(o)', 'abuelo': 'Abuela(o)',
    'abuela(o)': 'Abuela(o)', 'tia': 'Tia(o)', 'tio': 'Tia(o)',
    'tia(o)': 'Tia(o)', 'madre': 'Madre', 'padre': 'Padre',
    'madrastra': 'Madrastra', 'padrastro': 'Padrastro',
}

incidencias = collections.defaultdict(list)


# ------------------------------------------------------------------ utils --
def txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v == int(v):
        v = int(v)
    return re.sub(r'\s+', ' ', str(v)).strip()


def sin_tildes(s):
    return ''.join(c for c in unicodedata.normalize('NFD', s)
                   if unicodedata.category(c) != 'Mn').lower()


def sql(v):
    if v is None or v == '':
        return 'NULL'
    if isinstance(v, bool):
        return '1' if v else '0'
    if isinstance(v, int):
        return str(v)
    return "'" + str(v).replace('\\', '\\\\').replace("'", "''") + "'"


def fecha(d, m, a):
    d, m, a = txt(d), txt(m), txt(a)
    if not (d and m and a):
        return None
    if not m.isdigit():
        m = MESES.get(sin_tildes(m))
        if not m:
            return None
    try:
        d, m, a = int(d), int(m), int(a)
        if not (1 <= d <= 31 and 1 <= m <= 12 and 1900 < a < 2100):
            return None
        return '%04d-%02d-%02d' % (a, m, d)
    except ValueError:
        return None


def parse_grupo(valor):
    """'10-3' -> (10, 3) · '9o' -> (9, None) · 'No'/'' -> None."""
    s = txt(valor)
    if not s or s.lower() == 'no':
        return None
    m = re.match(r'^(\d{1,2})\s*-\s*(\d{1,2})$', s)
    if m:
        return int(m.group(1)), int(m.group(2))
    m = re.match(r'^(\d{1,2})\s*[ºo°]?$', s)
    if m:
        return int(m.group(1)), None
    return 'INVALIDO'


def norm_jornada(valor, donde):
    s = txt(valor)
    if not s:
        return None
    j = JORNADA_ALIAS.get(sin_tildes(s))
    if j:
        if sin_tildes(s) == 'trade':
            incidencias['jornada_typo'].append('%s · "Trade" -> Tarde' % donde)
        return j
    incidencias['jornada_sin_definir'].append('%s · "%s" -> NULL' % (donde, s))
    return None


# ---------------------------------------------------------------- acumula --
estudiantes = {}        # numero_documento -> dict  (la identidad es el numero)
acudientes = {}         # doc -> dict
barrios = {}            # clave sin tildes -> dict
grupos = {}             # (anio, sede_id, grado, numero, jornada) -> dict
matriculas = []
boletines_pendientes = []
boletines_crudos = []          # (matricula_id, numero, valor literal de la celda)
sede_discrepancia = []
grupo_discrepancia = []
resumen_libros = []


def id_barrio(nombre):
    """La base usa utf8mb4_unicode_ci: 'López' == 'Lopez' para el UNIQUE."""
    if not nombre:
        return None
    clave = sin_tildes(nombre)
    if clave not in barrios:
        barrios[clave] = {'id': len(barrios) + 1,
                          'formas': collections.Counter()}
    barrios[clave]['formas'][nombre.title()] += 1
    return barrios[clave]['id']


def id_grupo(anio, sede_id, grado, numero, jornada):
    clave = (anio, sede_id, grado, numero, jornada)
    if clave not in grupos:
        grupos[clave] = {
            'id': len(grupos) + 1, 'anio': anio, 'sede': sede_id,
            'grado': grado, 'numero': numero,
            'codigo': '%d-%d' % (grado, numero), 'jornada': jornada,
        }
    return grupos[clave]['id']


# ------------------------------------------------- pre-pasada: duplicados --
# 41 documentos aparecen en dos hojas o en dos sedes distintas (la misma
# persona registrada dos veces). En 35 casos la columna A desempata: solo
# una de las dos filas esta marcada '1' = matriculado activo. Se elige esa;
# cuando no hay forma de decidir, se toma la primera y se reporta.
def escanear_ocurrencias():
    ocur = collections.defaultdict(list)
    for idx, (cod_sede, archivo) in enumerate(LIBROS):
        ruta = os.path.join(DESCARGAS, archivo)
        if not os.path.exists(ruta):
            continue
        wb = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        for hoja in wb.sheetnames:
            if not re.match(r'^\d{1,2}-\d{1,2}$', hoja.strip()):
                continue
            ws = wb[hoja]
            for r, fila in enumerate(ws.iter_rows(min_row=3, max_col=40,
                                                  values_only=True), start=3):
                nombre = txt(fila[5]) if len(fila) > 5 else ''
                doc = txt(fila[39]) if len(fila) > 39 else ''
                if not nombre or not doc:
                    continue
                marca = txt(fila[0]).upper()
                ocur[doc].append((idx, cod_sede, hoja.strip(), r, marca, nombre))
        wb.close()
    return ocur


print('Resolviendo documentos repetidos entre libros ...')
ocurrencias = escanear_ocurrencias()
elegida = {}
duplicados_csv = []
for doc, ap in ocurrencias.items():
    if len(ap) == 1:
        elegida[doc] = (ap[0][0], ap[0][2], ap[0][3])
        continue
    activos = [a for a in ap if a[4] == '1']
    if len(activos) == 1:
        gana, motivo = activos[0], 'única fila marcada como activa'
    elif len(activos) > 1:
        gana, motivo = activos[0], 'AMBIGUO: varias filas activas, revisar'
        incidencias['duplicado_ambiguo'].append(
            '%s · %s · %s' % (doc, ap[0][5],
                              ' vs '.join('%s!%s' % (a[1], a[2]) for a in activos)))
    else:
        gana, motivo = ap[0], 'ninguna fila activa, se tomó la primera'
    elegida[doc] = (gana[0], gana[2], gana[3])
    for a in ap:
        duplicados_csv.append([
            doc, a[5], a[1], a[2], a[3], a[4] or '(vacío)',
            'CONSERVADA' if (a[0], a[2], a[3]) == elegida[doc] else 'omitida',
            motivo])
print('  %d documentos repetidos, %d filas involucradas'
      % (len([d for d, a in ocurrencias.items() if len(a) > 1]), len(duplicados_csv)))

# ------------------------------------------------------------------ carga --
for idx_libro, (cod_sede, archivo) in enumerate(LIBROS):
    ruta = os.path.join(DESCARGAS, archivo)
    if not os.path.exists(ruta):
        print('  !! falta %s' % archivo)
        incidencias['libro_faltante'].append(archivo)
        continue

    sede_id = SEDE[cod_sede]
    wb = openpyxl.load_workbook(ruta, data_only=True)
    hojas = [s for s in wb.sheetnames if re.match(r'^\d{1,2}-\d{1,2}$', s.strip())]
    print('%-4s %-38s %2d hojas' % (cod_sede, archivo[:38], len(hojas)), end='')

    n_libro = 0
    for hoja in hojas:
        ws = wb[hoja]
        etiqueta = hoja.strip()                     # 'CP' trae '4-1 '
        grado_hoja, numero_hoja = (int(x) for x in etiqueta.split('-'))

        # --- 1a pasada: juntar filas y votar la jornada del grupo -------
        filas = []
        votos = collections.Counter()
        for r in range(3, ws.max_row + 1):
            if not txt(ws.cell(row=r, column=6).value):
                continue
            filas.append(r)
            j = norm_jornada(ws.cell(row=r, column=33).value,
                             '%s!%s fila %d' % (cod_sede, etiqueta, r))
            if j:
                votos[j] += 1
        if not filas:
            continue
        jornada_grupo = votos.most_common(1)[0][0] if votos else 'Mañana'
        if len(votos) > 1:
            incidencias['hoja_jornada_mixta'].append(
                '%s!%s · %s -> el grupo queda como %s'
                % (cod_sede, etiqueta, dict(votos), jornada_grupo))

        # --- 2a pasada: procesar cada estudiante ------------------------
        for r in filas:
            C = lambda c: ws.cell(row=r, column=c).value
            nombre = txt(C(6))
            fila = '%s!%s fila %d' % (cod_sede, etiqueta, r)

            tipo = txt(C(39)) or 'T.I.'
            if tipo not in TIPO_DOC_OK:
                incidencias['tipo_doc_desconocido'].append(
                    '%s · %s · "%s" -> T.I.' % (fila, nombre, tipo))
                tipo = 'T.I.'
            doc = txt(C(40))
            if not doc:
                incidencias['sin_documento'].append('%s · %s' % (fila, nombre))
                continue

            # De las ocurrencias repetidas solo se procesa la elegida arriba.
            if elegida.get(doc) != (idx_libro, etiqueta, r):
                otro = elegida.get(doc)
                incidencias['documento_duplicado'].append(
                    '%s · %s · doc %s -> se conserva %s!%s fila %s'
                    % (fila, nombre, doc,
                       LIBROS[otro[0]][0] if otro else '?',
                       otro[1] if otro else '?', otro[2] if otro else '?'))
                continue

            est_id = len(estudiantes) + 1
            estudiantes[doc] = {
                'id': est_id, 'tipo': tipo, 'doc': doc, 'nombre': nombre,
                'nacimiento': fecha(C(35), C(36), C(37)),
                'genero': txt(C(34)) if txt(C(34)) in ('F', 'M') else None,
                'foto': txt(C(51)).upper() == 'S', 'origen': fila,
            }
            n_libro += 1

            # ---- acudiente (compartido entre sedes: hermanos) ----------
            acu_doc, acu_nombre = txt(C(53)), txt(C(52))
            if acu_doc and acu_nombre:
                if acu_doc not in acudientes:
                    acudientes[acu_doc] = {
                        'id': len(acudientes) + 1, 'doc': acu_doc,
                        'nombre': acu_nombre, 'fijo': txt(C(57)) or None,
                        'cel': txt(C(58)) or None, 'dir': txt(C(55)) or None,
                        'barrio_id': id_barrio(txt(C(56))), 'hijos': [],
                    }
                else:
                    id_barrio(txt(C(56)))      # cuenta la grafia del barrio
                par = PARENTESCO_ALIAS.get(sin_tildes(txt(C(54))))
                if par is None and txt(C(54)):
                    incidencias['parentesco_normalizado'].append(
                        '%s · "%s" -> Otro' % (fila, txt(C(54))))
                acudientes[acu_doc]['hijos'].append(
                    (est_id, PARENTESCO[par or 'Otro']))
            elif acu_nombre:
                incidencias['acudiente_sin_cedula'].append(
                    '%s · %s · "%s"' % (fila, nombre, acu_nombre))

            # ---- historial 2015..2025 desde los pares de columnas ------
            historial = {}
            for i, anio in enumerate(range(2015, 2027)):
                col_g = COL_ANIO_INICIO + i * 2
                g = parse_grupo(C(col_g))
                if g is None:
                    continue
                if g == 'INVALIDO':
                    incidencias['grupo_ilegible'].append(
                        '%s · %s · %d = "%s"' % (fila, nombre, anio, txt(C(col_g))))
                    continue
                cs = txt(C(col_g + 1)).upper()
                if cs not in SEDE:
                    if anio != ANIO_ACTUAL:
                        incidencias['sede_desconocida'].append(
                            '%s · %s · %d código "%s" -> omitido'
                            % (fila, nombre, anio, cs or '(vacío)'))
                        continue
                    cs = cod_sede
                historial[anio] = (g[0], g[1], SEDE[cs])

            # El ano en curso lo manda el libro y la hoja, no las columnas.
            ae = txt(C(31))
            af = txt(C(32)).upper()
            if ae != etiqueta:
                grupo_discrepancia.append(
                    [cod_sede, etiqueta, nombre, doc, ae or '(vacío)',
                     txt(C(29)), etiqueta])
            if af and af != cod_sede:
                sede_discrepancia.append(
                    [cod_sede, etiqueta, nombre, doc, 'AF=' + af,
                     SEDE_NOMBRE[cod_sede]])
            elif txt(C(7)) and txt(C(7)) != SEDE_NOMBRE[cod_sede]:
                sede_discrepancia.append(
                    [cod_sede, etiqueta, nombre, doc, 'SEDE=' + txt(C(7)),
                     SEDE_NOMBRE[cod_sede]])
            historial[ANIO_ACTUAL] = (grado_hoja, numero_hoja, sede_id)

            # ---- una matricula por ano cursado ------------------------
            for anio in sorted(historial):
                grado, numero, sid = historial[anio]
                actual = anio == ANIO_ACTUAL
                jornada_est = norm_jornada(C(33), fila) if actual else None

                if numero is not None:
                    gid = id_grupo(anio, sid, grado, numero,
                                   jornada_grupo if actual else 'Mañana')
                else:
                    gid = None
                    incidencias['grupo_sin_numero'].append(
                        '%s · %s · %d grado %d sin número de grupo'
                        % (fila, nombre, anio, grado))

                prev = historial.get(anio - 1)
                cond = ('nuevo' if prev is None
                        else 'repitente' if prev[0] == grado else 'antiguo')
                sig = historial.get(anio + 1)
                resultado = None
                if sig:
                    resultado = 'promovido' if sig[0] > grado else 'reprobado'

                estado = 'activo'
                if actual:
                    marca = txt(C(1)).upper()
                    if marca == 'R':
                        estado = 'retirado'
                    elif marca != '1':
                        estado = 'cancelado'
                        incidencias['sin_marca_matricula'].append(
                            '%s · %s · columna A = "%s" -> cancelado'
                            % (fila, nombre, marca or '(vacío)'))

                modalidad = None
                if actual and grado in GRADOS_MEDIA:
                    mod = txt(C(8))
                    modalidad = MODALIDAD.get(mod)
                    if mod and modalidad is None:
                        incidencias['modalidad_desconocida'].append(
                            '%s · "%s"' % (fila, mod))
                elif actual and txt(C(8)):
                    incidencias['modalidad_fuera_de_la_media'].append(
                        '%s · grado %d trae "%s" -> ignorado'
                        % (fila, grado, txt(C(8))))

                matriculas.append({
                    'id': len(matriculas) + 1, 'est': est_id,
                    'anio': ANIO[anio], 'grado': GRADO[grado], 'grupo': gid,
                    'sede': sid, 'modalidad': modalidad,
                    'jornada': jornada_est, 'cond': cond, 'estado': estado,
                    'resultado': resultado, 'historico': not actual,
                    'fecha': fecha(C(3), C(4), C(5)) if actual else None,
                    'orden': int(txt(C(2))) if actual and txt(C(2)).isdigit() else None,
                    'obs': (txt(C(59)) or None) if actual else None,
                })

            # Las columnas BOLETIN se guardan literales en `boletines_excel`:
            # no se sabe si sus numeros son periodos del ano o grados, asi que
            # no se traducen a `entregas_boletin` todavia. La matricula del ano
            # en curso es siempre la ultima que se agrego para este estudiante.
            bol = {}
            for i, c in enumerate(COL_BOLETIN, start=1):
                if txt(C(c)):
                    bol[i] = txt(C(c))
            if bol:
                assert matriculas[-1]['anio'] == ANIO[ANIO_ACTUAL]
                mat_id = matriculas[-1]['id']
                for numero, valor in sorted(bol.items()):
                    boletines_crudos.append((mat_id, numero, valor))
                    if valor.upper() not in ('S', 'NO APLICA'):
                        incidencias['boletin_valor_raro'].append(
                            '%s · %s · boletín %dº = "%s"'
                            % (fila, nombre, numero, valor))
                boletines_pendientes.append(
                    [cod_sede, etiqueta, nombre, doc]
                    + [bol.get(i, '') for i in range(1, 11)])

    print('  -> %4d estudiantes' % n_libro)
    resumen_libros.append((cod_sede, SEDE_NOMBRE[cod_sede], len(hojas), n_libro))
    wb.close()

print('\nTOTALES  estudiantes=%d  acudientes=%d  barrios=%d  grupos=%d  '
      'matrículas=%d  boletines=%d'
      % (len(estudiantes), len(acudientes), len(barrios), len(grupos),
         len(matriculas), len(boletines_crudos)))

# ------------------------------------------------------------------- SQL ---
out = io.open(os.path.join(OUT, 'datos-sedes.sql'), 'w',
              encoding='utf-8', newline='\n')
w = out.write
w('-- Generado por docs/bd/importar_sedes.py\n')
w('-- Carga completa de los %d libros de sede sobre la base `sieage`.\n' % len(LIBROS))
w('-- Requiere que las tablas de datos esten vacias y los catalogos sembrados.\n')
w('SET NAMES utf8mb4;\nSTART TRANSACTION;\n\n')

w('-- barrios ------------------------------------------------------------\n')
for b in sorted(barrios.values(), key=lambda x: x['id']):
    nombre = b['formas'].most_common(1)[0][0]
    if len(b['formas']) > 1:
        incidencias['barrio_unificado'].append(
            '%s <- %s' % (nombre, ', '.join(sorted(b['formas']))))
    w('INSERT INTO barrios (id, nombre) VALUES (%d, %s);\n' % (b['id'], sql(nombre)))

w('\n-- grupos ------------------------------------------------------------\n')
for g in sorted(grupos.values(), key=lambda x: x['id']):
    w('INSERT INTO grupos (id, anio_lectivo_id, sede_id, grado_id, numero, '
      'codigo, jornada, cupos_proyectados) VALUES (%d, %d, %d, %d, %d, %s, %s, 34);\n'
      % (g['id'], ANIO[g['anio']], g['sede'], GRADO[g['grado']], g['numero'],
         sql(g['codigo']), sql(g['jornada'])))

w('\n-- estudiantes -------------------------------------------------------\n')
for e in sorted(estudiantes.values(), key=lambda x: x['id']):
    w('INSERT INTO estudiantes (id, tipo_documento, numero_documento, '
      'nombre_completo, fecha_nacimiento, genero, tiene_foto) VALUES '
      '(%d, %s, %s, %s, %s, %s, %s);\n'
      % (e['id'], sql(e['tipo']), sql(e['doc']), sql(e['nombre']),
         sql(e['nacimiento']), sql(e['genero']), sql(e['foto'])))

w('\n-- acudientes --------------------------------------------------------\n')
for a in sorted(acudientes.values(), key=lambda x: x['id']):
    w('INSERT INTO acudientes (id, tipo_documento, numero_documento, '
      'nombre_completo, telefono_fijo, telefono_celular, direccion, barrio_id) '
      "VALUES (%d, 'C.C.', %s, %s, %s, %s, %s, %s);\n"
      % (a['id'], sql(a['doc']), sql(a['nombre']), sql(a['fijo']),
         sql(a['cel']), sql(a['dir']), sql(a['barrio_id'])))

w('\n-- vinculo estudiante-acudiente --------------------------------------\n')
for a in sorted(acudientes.values(), key=lambda x: x['id']):
    for est_id, par_id in a['hijos']:
        w('INSERT INTO estudiante_acudiente (estudiante_id, acudiente_id, '
          'parentesco_id) VALUES (%d, %d, %d);\n' % (est_id, a['id'], par_id))

w('\n-- matriculas (una fila por ano cursado) -----------------------------\n')
for m in matriculas:
    w('INSERT INTO matriculas (id, estudiante_id, anio_lectivo_id, grado_id, '
      'grupo_id, sede_id, modalidad_id, jornada, fecha_matricula, condicion, '
      'estado, resultado, numero_orden, es_historico, observaciones) VALUES '
      '(%d, %d, %d, %d, %s, %d, %s, %s, %s, %s, %s, %s, %s, %s, %s);\n'
      % (m['id'], m['est'], m['anio'], m['grado'], sql(m['grupo']), m['sede'],
         sql(m['modalidad']), sql(m['jornada']), sql(m['fecha']),
         sql(m['cond']), sql(m['estado']), sql(m['resultado']),
         sql(m['orden']), sql(m['historico']), sql(m['obs'])))

w('\n-- boletines: columnas BOLETIN del Excel, literales y sin interpretar ---\n')
for mat_id, numero, valor in boletines_crudos:
    w('INSERT INTO boletines_excel (matricula_id, numero, valor) VALUES (%d, %d, %s);\n'
      % (mat_id, numero, sql(valor)))

w('\nCOMMIT;\n')
out.close()


# --------------------------------------------------------------- reportes --
def csv_out(nombre, cabecera, filas):
    if not filas:
        return
    with io.open(os.path.join(OUT, nombre), 'w', encoding='utf-8-sig',
                 newline='') as f:
        cw = csv.writer(f, delimiter=';')
        cw.writerow(cabecera)
        cw.writerows(filas)


csv_out('pendiente-documentos-duplicados.csv',
        ['documento', 'estudiante', 'sede', 'grupo', 'fila', 'columna_A',
         'decision', 'motivo'], duplicados_csv)
csv_out('pendiente-boletines.csv',
        ['sede', 'grupo', 'estudiante', 'documento']
        + ['boletin_%d' % i for i in range(1, 11)], boletines_pendientes)
csv_out('pendiente-grupo-desactualizado.csv',
        ['sede', 'hoja', 'estudiante', 'documento', 'columna_GRUPO_2026',
         'columna_GRUPO_2025', 'grupo_importado'], grupo_discrepancia)
csv_out('pendiente-sede-discrepancia.csv',
        ['sede_libro', 'grupo', 'estudiante', 'documento', 'valor_en_el_excel',
         'sede_importada'], sede_discrepancia)

rep = io.open(os.path.join(OUT, 'reporte-sedes.txt'), 'w',
              encoding='utf-8', newline='\n')
rep.write('REPORTE DE IMPORTACIÓN MULTI-SEDE · matrícula 2026\n')
rep.write('=' * 72 + '\n\nLIBROS PROCESADOS\n')
for cod, nom, nh, ne in resumen_libros:
    rep.write('  %-3s %-22s %2d grupos  %4d estudiantes\n' % (cod, nom, nh, ne))
rep.write('  %-26s %2d grupos  %4d estudiantes\n'
          % ('TOTAL', sum(x[2] for x in resumen_libros),
             sum(x[3] for x in resumen_libros)))

rep.write('\nREGISTROS GENERADOS\n')
for k, v in [('estudiantes', len(estudiantes)), ('acudientes', len(acudientes)),
             ('barrios', len(barrios)), ('grupos', len(grupos)),
             ('matrículas', len(matriculas)),
             ('boletines (crudos)', len(boletines_crudos))]:
    rep.write('  %-14s %d\n' % (k, v))

inv = {v: k for k, v in ANIO.items()}
rep.write('\nMATRÍCULAS POR AÑO\n')
for aid, n in sorted(collections.Counter(m['anio'] for m in matriculas).items()):
    rep.write('  %d  %5d\n' % (inv[aid], n))

rep.write('\nAÑO 2026 POR SEDE Y ESTADO\n')
c = collections.Counter((m['sede'], m['estado']) for m in matriculas
                        if m['anio'] == ANIO[ANIO_ACTUAL])
inv_sede = {v: k for k, v in SEDE.items()}
for (sid, est), n in sorted(c.items()):
    rep.write('  %-22s %-10s %4d\n' % (SEDE_NOMBRE[inv_sede[sid]], est, n))

rep.write('\nINCIDENCIAS\n')
if not incidencias:
    rep.write('  ninguna\n')
for k in sorted(incidencias):
    rep.write('\n  [%s]  %d caso(s)\n' % (k, len(incidencias[k])))
    for linea in incidencias[k][:30]:
        rep.write('    - %s\n' % linea)
    if len(incidencias[k]) > 30:
        rep.write('    ... y %d más\n' % (len(incidencias[k]) - 30))
rep.close()

print('\nGenerado:\n  docs/bd/datos-sedes.sql\n  docs/bd/reporte-sedes.txt')
for n, f in [('pendiente-boletines.csv', boletines_pendientes),
             ('pendiente-grupo-desactualizado.csv', grupo_discrepancia),
             ('pendiente-sede-discrepancia.csv', sede_discrepancia)]:
    if f:
        print('  docs/bd/%s (%d filas)' % (n, len(f)))
